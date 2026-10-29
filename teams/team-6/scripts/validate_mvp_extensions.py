"""Read-only artifact and cloud checks for slides, proximity mapping and monitoring."""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import json
import os
from pathlib import Path
from zipfile import ZipFile

import pyarrow.parquet as pq
from databricks.sdk import WorkspaceClient
from databricks.sdk.service import jobs
from citibike_pipeline.forecast_sources import digest


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--package',type=Path,required=True)
    parser.add_argument('--output',type=Path,default=Path('docs/validation/mvp_extensions_2026-10-16.json'))
    args=parser.parse_args()
    root=Path(__file__).resolve().parents[1]
    manifest=json.loads((args.package/'manifest.json').read_text())
    tables={}
    for artifact in manifest['artifacts']:
        assert digest(args.package/artifact['path'])==artifact['sha256']
        if artifact['path'].endswith('.parquet'):
            rows=pq.read_table(args.package/artifact['path']).to_pylist()
            assert len(rows)==artifact['rows'];tables[artifact['name']]=rows
    for name,key in (('station_proximity','gtfs_stop_id'),('complex_proximity','mta_complex_id'),('route_proximity','route_id')):
        rows=tables[name]
        keys={(r['bike_station_id'],r['radius_m'],r[key]) for r in rows}
        assert len(keys)==len(rows),name+' duplicate keys'
        assert all(0<=r['distance_m']<=r['radius_m'] for r in rows)
        smaller={(r['bike_station_id'],r[key]) for r in rows if r['radius_m']==250}
        larger={(r['bike_station_id'],r[key]) for r in rows if r['radius_m']==500}
        assert smaller<=larger
    assert len(tables['coverage'])==2*len(tables['bike_stations'])
    assert all(not r['historical_validity_verified'] and not r['station_operation_verified']
        for rows in tables.values() for r in rows)
    for rows in tables.values():
        assert all(isinstance(r['bike_station_id'],str) for r in rows)
    assert len({(r['bike_station_id'],r['radius_m']) for r in tables['coverage']})==len(tables['coverage'])
    os.environ['PATH']=str(root/'.local/bin')+os.pathsep+os.environ['PATH']
    w=WorkspaceClient(profile='DATA622')
    states={}
    for name,filename in (('health','databricks_health'),('mapping','borough_mapping'),('dashboard','borough_dashboard')):
        state=json.loads((root/'.local/execution'/(filename+'.json')).read_text())
        run=w.jobs.get_run(state['run_id'])
        assert run.state.result_state==jobs.RunResultState.SUCCESS,(name,run.state)
        result=w.jobs.get_run_output(run.tasks[0].run_id)
        output=result.notebook_output.result
        report=json.loads(output) if name!='dashboard' else output
        states[name]={'run_id':run.run_id,'state':run.state.as_dict(),'report':report}
        if name=='health':
            schedule=w.jobs.get(state['job_id']).settings.schedule
            assert schedule.pause_status==jobs.PauseStatus.UNPAUSED
            assert schedule.quartz_cron_expression=='0 15 13 * * ?' and schedule.timezone_id=='UTC'
            states[name]['job_id']=state['job_id'];states[name]['schedule']=schedule.as_dict()
    assert states['health']['report']['status']=='OK'
    assert states['mapping']['report']['mapping_version']==manifest['mapping_version']
    for name,rows in tables.items():
        assert states['mapping']['report']['tables'][name]['rows']==len(rows)
    # A successful notebook run also queried both newly created stable views.
    assert states['dashboard']['report']=='Manhattan forecast dashboard queries succeeded'
    slides=root/'docs/presentations'
    from pptx import Presentation
    import pymupdf
    deck=Presentation(slides/'data622_mvp.pptx')
    assert len(deck.slides)==3 and all(s.has_notes_slide for s in deck.slides)
    assert all(s.notes_slide.notes_text_frame.text.strip() for s in deck.slides)
    assert len(pymupdf.open(slides/'data622_mvp.pdf'))==3
    with ZipFile(slides/'data622_mvp.pptx') as z:assert z.testzip() is None
    output={'recorded_at_utc':datetime.now(timezone.utc).isoformat(),
        'mapping_package':str(args.package),'mapping_manifest_sha256':digest(args.package/'manifest.json'),
        'mapping_checks':['all artifact SHA-256 values match','unique stop/complex/route keys within radius',
            'distances within configured radius','250m keys subset of 500m keys','all bike/radius coverage rows retained',
            'IDs are strings','historical validity and station operation remain unverified','cloud table counts match'],
        'coverage':manifest['coverage'],'local_table_rows':{k:len(v) for k,v in tables.items()},**states,
        'slides':{'count':3,'speaker_notes':True,'pdf_pages':3,'powerpoint_zip_valid':True}}
    args.output.write_text(json.dumps(output,indent=2)+'\n')
    print(json.dumps({k:v for k,v in output.items() if k not in ('health','mapping')},indent=2))


if __name__=='__main__':main()
