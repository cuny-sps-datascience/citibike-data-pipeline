"""Render the three-slide DATA622 demo as editable PowerPoint and matching PDF.

All text, diagrams and bars are native shapes. No training or live forecast job
is triggered. Evaluation numbers come from recorded cloud validation evidence.
"""
from __future__ import annotations

import argparse
from datetime import date
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

ROOT=Path(__file__).resolve().parents[1]
W,H=13.333333,7.5
NAVY='102A43'; TEAL='007F7B'; MINT='DDF4EB'; INK='15334B'
GRAY='526577'; PALE='F1F5F8'; WHITE='FFFFFF'; AMBER='A3530A'; SAND='FFF1DB'


class Deck:
    def __init__(self, destination):
        destination.mkdir(parents=True,exist_ok=True)
        self.destination=destination
        self.prs=Presentation();self.prs.slide_width=Inches(W);self.prs.slide_height=Inches(H)
        self.prs.core_properties.title='Citi Bike: three-hour Manhattan forecast MVP'
        self.prs.core_properties.author='Randy Howk'
        self.pdf=canvas.Canvas(str(destination/'data622_mvp.pdf'),pagesize=(W*72,H*72))
        self.pdf.setTitle(self.prs.core_properties.title);self.pdf.setAuthor('Randy Howk')

    def slide(self, number, section, title, subtitle):
        if number>1:self.pdf.showPage()
        self.current=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(0,0,W,H,WHITE)
        self.rect(.48,.43,.07,.24,TEAL)
        self.text(.7,.42,11,.32,'DATA 622  /  RANDY HOWK  /  '+section,11,TEAL,True)
        self.text(.65,.95,12,.65,title,31,NAVY,True)
        self.text(.67,1.68,12,.48,subtitle,14,GRAY)
        self.line(.65,7.06,12.68,7.06,'D7E1E9',.7)
        self.text(.67,7.17,11,.2,'PUBLIC-DATA MVP  |  Prepared '+date.today().isoformat()+'  |  Manhattan totals',9,GRAY)
        self.text(12.15,7.13,.5,.3,f'{number:02d}',12,TEAL,True)

    def rect(self,x,y,w,h,color,radius=False):
        shape=self.current.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Inches(x),Inches(y),Inches(w),Inches(h))
        if radius:shape.adjustments[0]=.08
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color);shape.line.fill.background()
        self.pdf.setFillColor('#'+color)
        if radius:self.pdf.roundRect(x*72,(H-y-h)*72,w*72,h*72,7,stroke=0,fill=1)
        else:self.pdf.rect(x*72,(H-y-h)*72,w*72,h*72,stroke=0,fill=1)

    def text(self,x,y,w,h,value,size=18,color=INK,bold=False,link=None):
        box=self.current.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        frame=box.text_frame;frame.clear();frame.word_wrap=False
        frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
        self.pdf.setFillColor('#'+color);self.pdf.setFont('Helvetica-Bold' if bold else 'Helvetica',size)
        for index,line in enumerate(value.split('\n')):
            p=frame.paragraphs[0] if index==0 else frame.add_paragraph()
            p.space_after=Pt(0);p.space_before=Pt(0);p.line_spacing=Pt(size*1.22)
            run=p.add_run();run.text=line
            run.font.name='Arial';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
            if link:run.hyperlink.address=link
            self.pdf.drawString(x*72,(H-y)*72-size*.93-index*size*1.22,line)
            if self.pdf.stringWidth(line,'Helvetica-Bold' if bold else 'Helvetica',size)>w*72+1:
                raise ValueError('Slide text exceeds assigned width: '+line)
        if len(value.split('\n'))*size*1.22>h*72+4:
            raise ValueError('Slide text exceeds assigned height: '+value)
        if link:self.pdf.linkURL(link,(x*72,(H-y-h)*72,(x+w)*72,(H-y)*72),relative=0)

    def line(self,x1,y1,x2,y2,color=TEAL,width=2):
        from pptx.enum.shapes import MSO_CONNECTOR
        shape=self.current.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        shape.line.color.rgb=RGBColor.from_string(color);shape.line.width=Pt(width)
        self.pdf.setStrokeColor('#'+color);self.pdf.setLineWidth(width)
        self.pdf.line(x1*72,(H-y1)*72,x2*72,(H-y2)*72)

    def down(self,x,y1,y2):
        self.line(x,y1,x,y2);self.line(x,y2,x-.07,y2-.07);self.line(x,y2,x+.07,y2-.07)

    def image(self,path,x,y,w,h):
        self.current.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
        self.pdf.drawImage(ImageReader(str(path)),x*72,(H-y-h)*72,w*72,h*72)

    def notes(self,text):
        self.current.notes_slide.notes_text_frame.text=text

    def save(self):
        self.prs.save(self.destination/'data622_mvp.pptx');self.pdf.save()


def build(destination):
    metrics=json.loads((ROOT/'docs/validation/public_only_mvp_2026-10-15.json').read_text())['training']
    model=metrics['selected_model'];test=metrics['test'];selected=test[model]['mae']
    health_path=ROOT/'docs/validation/september_health_2026-10-16.json'
    if not health_path.exists():health_path=ROOT/'docs/validation/mvp_extensions_2026-10-16.json'
    health=json.loads(health_path.read_text())['health']['report'] if health_path.exists() else None
    rolling=json.loads((ROOT/'docs/validation/release_aware_evidence_2026-10-16.json').read_text())['runs']['rolling']['outputs']['review']
    selected_rolling=rolling['selected_pooled_metrics']['mae']
    profile_rolling=rolling['profile_pooled_metrics']['mae']
    deck=Deck(destination)
    deck.slide(1,'PROBLEM + SCOPE','Citi Bike departures, three hours ahead',
        'A working Manhattan-total forecast using public trips and forecast weather in Databricks Free Edition.')
    deck.rect(.65,2.38,5.55,2.25,NAVY,True)
    deck.text(.92,2.65,4.9,.4,'Plan for the next three hours',23,WHITE,True)
    deck.text(.92,3.31,4.95,1.04,'Estimate served departures across Manhattan.\nPublish before the first target hour.\nCompare weather with a calendar fallback.',17,WHITE)
    for i,(label,interval) in enumerate((('HOUR 1','T to T+1'),('HOUR 2','T+1 to T+2'),('HOUR 3','T+2 to T+3'))):
        x=6.55+i*2.05
        deck.rect(x,2.72,1.88,1.48,MINT,True)
        deck.text(x+.15,2.97,1.6,.3,label,13,TEAL,True)
        deck.text(x+.15,3.53,1.6,.35,interval,14,INK,True)
    deck.text(6.6,4.36,5.65,.35,'Three separate hourly totals at each origin T',15,GRAY)
    deck.rect(.65,5.03,5.55,1.56,PALE,True)
    deck.text(.88,5.23,5.05,.32,'IMPLEMENTED',12,TEAL,True)
    deck.text(.88,5.69,5.05,.65,'Hourly forecasts + dashboard + daily health\nDaily releases, outcomes and station evidence',16)
    deck.rect(6.55,5.03,6.1,1.56,SAND,True)
    deck.text(6.79,5.23,5.55,.32,'WHY START WITH MANHATTAN?',12,AMBER,True)
    deck.text(6.79,5.69,5.55,.65,'Station archive coverage depends on activity.\nReliable operation history is still needed.',16)
    deck.notes('Timing: 90 seconds. '+
        'The original proposal predicts station departures using weather and subway disruptions. This public-data MVP reduces spatial scope to Manhattan totals; it does not complete the original proposal. '+
        'Labels use each trip start coordinate within Manhattan. These are served trips, not latent or unmet demand. '+
        'The CityBikes public archive records inventory changes, rather than continuous station snapshots. Only 0.147% of its conservatively covered station hours have zero departures; that selection cannot support unconditional station forecasting. '+
        'Each output is one separate hourly count, not a cumulative three-hour count. '+
        'Course format: https://georgehagstrom.github.io/DATA622Fall2026/assignments/project.html')

    deck.slide(2,'SYSTEM','From public data to a persisted forecast',
        'Bulk trips stay on the RAID. Compact features, models, forecasts and health reports live in Databricks.')
    deck.text(.85,2.25,5.5,.3,'BATCH DEVELOPMENT',12,TEAL,True)
    deck.text(7.25,2.25,5.1,.3,'HOURLY SERVING  /  :55 UTC',12,TEAL,True)
    for x,labels in ((.7,[('Public trips + archived weather','Geographic labels; released trip profiles'),
                         ('Chronological model comparison','Initial selection + separate monthly release audit'),
                         ('Frozen model + MLflow report','Weather Poisson + calendar fallback')]),
                     (7.1,[('Collect live public inputs','Forecast weather; GBFS; MTA alerts archived'),
                           ('Score the next three hours','Weather when usable; calendar otherwise'),
                           ('Persist -> dashboard + daily health','Preserve first prediction; evaluate released outcomes')])):
        for j,(title,body) in enumerate(labels):
            y=2.76+j*1.05
            deck.rect(x,y,5.5,.79,PALE if x<1 else MINT,True)
            deck.text(x+.18,y+.11,5.12,.33,title,18,INK,True)
            deck.text(x+.18,y+.48,5.12,.25,body,11,GRAY)
            if j<2:deck.down(x+2.75,y+.79,y+1.05)
    deck.line(6.2,5.25,6.65,5.25)
    deck.line(6.65,5.25,6.65,4.2)
    deck.line(6.65,4.2,7.1,4.2)
    deck.line(7.1,4.2,7.02,4.13);deck.line(7.1,4.2,7.02,4.27)
    deck.rect(.7,6.1,11.9,.59,NAVY,True)
    deck.text(.9,6.26,11.5,.29,'Input cutoff T-5 min  |  Trip release +45 days*  |  Weather run +7 hours*  |  First publication wins',14,WHITE)
    deck.text(.73,6.79,11.9,.2,'*Availability assumptions; early archived weather includes hindcasts. GBFS/MTA collection is not used by the current borough model.',9,GRAY)
    deck.notes('Timing: 100 seconds. '+
        'Training March 2025-March 2026, validation April-May 2026, test June-August 2026. Purge target intervals crossing split boundaries; ambiguous fall-back labels remain null. '+
        'No recent trip lags or future trip properties are predictors. Forecast inputs use a five-minute cutoff, assumed 45-day monthly trip release and seven-hour weather run availability. Historical provider publication times are not established; early weather archives include hindcasts. '+
        'The additional monthly evaluation gates fitting, validation selection and refitting on assumed label release; June selects March validation, July and August select April. Earlier inspected test months remain retrospective evidence. '+
        'Live collection and scoring run hourly at :55 UTC with one concurrent run and bounded timeouts. Calendar fallback covers stale/missing weather. '+
        'Model artifacts and input packages are checked by SHA-256. A late retry preserved the first forecasts and publication times. '+
        'Daily health measures forecast availability and source retrieval freshness; it cannot establish predictive accuracy. '+
        'Daily evidence watches release versions, imports prepared validated outcomes and normalizes prospective station information and subway advisories. September data extends outcomes without retraining or activating a different model. '+
        'Live dashboard: https://dbc-d00cbd4c-c74c.cloud.databricks.com/editor/notebooks/2763628353985993')

    deck.slide(3,'EVIDENCE + NEXT STEPS','A useful benchmark, with clear limits',
        f'June-August 2026 release-aware backtests: {rolling["selected_pooled_metrics"]["rows"]:,} common forecast rows across three horizons.')
    deck.text(.75,2.33,5.5,.33,'MEAN ABSOLUTE ERROR  /  departures per hour',12,GRAY,True)
    series=[('Hour-of-week profile',profile_rolling,'91A3B2'),
            ('Monthly selected weather model',selected_rolling,TEAL)]
    for index,(name,value,color) in enumerate(series):
        y=3.02+index*1.03
        deck.text(.77,y,4.7,.28,name,15,INK,index==1)
        width=4.25*value/profile_rolling
        deck.rect(.78,y+.34,width,.22,color,True)
        deck.text(.9+width,y+.28,.8,.35,f'{value:,.0f}',15,color,True)
    deck.rect(.73,5.34,5.4,.77,MINT,True)
    improvement=(1-selected_rolling/profile_rolling)*100
    deck.text(.94,5.54,4.98,.38,f'{improvement:.1f}% lower MAE than the weekly profile',18,TEAL,True)
    deck.text(6.6,2.36,5.9,.35,'Monthly selected model MAE',15,INK,True)
    for index,fold in enumerate(rolling['folds']):
        y=2.98+index*.46
        deck.text(6.65,y,3.7,.3,fold['month'],15,GRAY)
        deck.text(10.8,y,1.3,.3,f'{fold["test"][fold["selected_model"]]["mae"]:,.0f}',15,TEAL,True)
    deck.text(6.65,4.47,5.65,.53,'June calendar model was better: 650 vs. 734.\nWeather does not improve every month.',13,AMBER)
    deck.text(6.65,5.14,5.6,.32,'REMAINING FOR THE ORIGINAL PROPOSAL',11,AMBER,True)
    deck.text(6.65,5.55,5.68,.89,'Verified station history + historical alert evidence\nClass deployment + documented shared catalog\nPublished live outcomes for production accuracy',14)
    deck.text(.78,6.59,10.1,.23,'Retrospective evidence with release/hindcast assumptions. Active model frozen; live accuracy pending.',10,GRAY)
    deck.text(11.1,6.58,1.48,.3,'Open live demo',12,TEAL,True,
        link='https://dbc-d00cbd4c-c74c.cloud.databricks.com/editor/notebooks/2763628353985993')
    deck.notes('Timing: 110 seconds. '+
        'Release-aware monthly selection gives pooled MAE 767.77 versus the weekly profile 1386.49 on 6552 common rows. June chooses weather alpha 1.0; July and August choose alpha 0.1. Monthly MAEs are 734.30, 839.87 and 727.24. June calendar comparison is better at 649.80. '+
        'The original frozen model test remains 761.00 on 6564 rows; it is a different design and the new analysis does not promote a model. The extra monthly-boundary purge explains the row-count difference. These months were already inspected, so this is not a new independent holdout. '+
        'The comparison pools overlapping targets across three horizons; rows are not independent. Calendar and weather comparisons use the same covered keys. '+
        'This is retrospective accuracy under availability and hindcast assumptions, not measured live accuracy or a causal effect. '+
        'Current 250m/500m station-to-subway mapping is dated from retrieval and cannot be backfilled into historical training. Straight-line proximity is not walking distance; reference daytime route names are not proof of disrupted service. '+
        'Reliable operation history, historical alerts and class deployment/sharing remain open. No production MAE is available until genuine trip outcomes mature. '+
        'Live demo, five minutes: 1) dashboard, horizons and publication timing; 2) model comparison and plot; 3) hourly job graph/latest success; 4) mapping coverage and weather handoff; 5) daily health and empty live evaluation. '+
        'Before presenting, refresh dashboard and checks without retraining. Full notes: docs/public_mvp_demo.md. Metrics: docs/manhattan_benchmark_model_card.md. '+
        'September outcomes and weather are prepared as a separate data refresh. They cannot evaluate October live forecasts. Timed navigation: docs/midterm_demo_script.md. '+
        'Public station sources: https://gbfs.citibikenyc.com/gbfs/gbfs.json and https://data.ny.gov/Transportation/MTA-Subway-Stations/39hk-dx4f/about_data')
    deck.save()
    import pymupdf
    pdf=pymupdf.open(destination/'data622_mvp.pdf')
    assert len(pdf)==3
    for index,page in enumerate(pdf):
        page.get_pixmap(matrix=pymupdf.Matrix(1.35,1.35)).save(destination/f'slide_{index+1}.png')
    (destination/'build_manifest.json').write_text(json.dumps({'slides':3,'source_notes':'docs/public_mvp_demo.md',
        'model_version':metrics['model_version'],'selected_model':model,
        'test_mae':selected,'test_rows':metrics['split_rows']['test'],
        'primary_evaluation':'release-aware monthly backtests; previously inspected retrospective months',
        'rolling_mae':selected_rolling,'rolling_profile_mae':profile_rolling,
        'rolling_rows':rolling['selected_pooled_metrics']['rows'],
        'health_snapshot_utc':health['checked_at_utc'] if health else None,
        'outputs':['data622_mvp.pptx','data622_mvp.pdf','slide_1.png','slide_2.png','slide_3.png']},indent=2)+'\n')
    print(destination)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=ROOT/'docs/presentations')
    build(parser.parse_args().output)
